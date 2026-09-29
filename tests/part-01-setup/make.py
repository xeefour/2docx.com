# -*- coding: utf-8 -*-
"""
ส่วนที่ 1: ติดตั้งและตรวจสอบระบบ
=================================
ทดสอบว่าระบบพร้อมใช้งานหรือไม่ ก่อนทดสอบส่วนอื่น

  1.1  Carbone ตอบเวอร์ชันได้
  1.2  ฟอนต์ TH Sarabun New อยู่ใน container
  1.3  ฟอนต์ไทยมาตรฐานอื่นอยู่ใน container
  1.4  ใช้หน่วยความจำไม่เกิน 1 GB
  1.5  แปลง DOCX เปล่าเป็น PDF ได้ (พื้นฐานของทุกอย่าง)

รัน:  python make.py
"""
import os
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
TPL_DIR = os.path.join(HERE, "templates")
DATA_DIR = os.path.join(HERE, "data")
os.makedirs(TPL_DIR, exist_ok=True)
os.makedirs(DATA_DIR, exist_ok=True)

FONT = "TH Sarabun New"
OUT = os.path.join(TPL_DIR, "ว่างเปล่า.docx")

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)

st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(16)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

p = doc.add_paragraph()
r = p.add_run("ทดสอบการแปลงไฟล์เปล่าเป็น PDF")
r.font.name = FONT
r.font.size = Pt(16)
r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

doc.save(OUT)
print("สร้างแม่แบบว่างเปล่า:", os.path.basename(OUT))

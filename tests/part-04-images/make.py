# -*- coding: utf-8 -*-
"""
ส่วนที่ 4: รูปภาพ
====================
Carbone เวอร์ชันฟรีใส่รูปเองไม่ได้ (เป็น Enterprise Feature)
ส่วนนี้จึงทดสอบ "ตัวฉีดรูปที่เขียนเอง" แล้วให้ Carbone แปลงเป็น PDF

  4.1  รูปเดี่ยว
  4.2  หลายรูปในช่องเดียว
  4.3  รูปในตาราง
  4.4  ซ่อนแถวที่ไม่มีรูป
  4.5  ลบรูปที่ไม่ต้องการ
  4.6  รูปแยกส่วนกันจริง (checksum ต่างกัน)
  4.7  แปลงเป็น PDF ได้

วิธีเตรียมแม่แบบ:
  คลิกขวารูป → คุณสมบัติ → Alt Text → ใส่ชื่อช่อง
  ตัวฉีดจับคู่จากช่องนี้

รัน:  python make.py
"""
import os, base64, json
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

HERE = os.path.dirname(os.path.abspath(__file__))
TPL = os.path.join(HERE, "templates", "รูปภาพ.docx")
DATA = os.path.join(HERE, "data", "ข้อมูล.json")
IMGD = os.path.join(HERE, "images")
FONT_DIR = os.path.join(HERE, "..", "lib", "fonts")
for d in (os.path.dirname(TPL), os.path.dirname(DATA), IMGD):
    os.makedirs(d, exist_ok=True)

FONT = "TH Sarabun New"


def thai(size, bold=False):
    f = "THSarabunNew Bold.ttf" if bold else "THSarabunNew.ttf"
    try:
        return ImageFont.truetype(os.path.join(FONT_DIR, f), size)
    except Exception:
        return ImageFont.load_default()


def make_img(path, color, label, size=(480, 340), fmt="PNG"):
    """สร้างรูปทดสอบ — ใช้สีและเส้นทแยง เพื่อให้แยกรูปออกจากกันง่าย"""
    W, H = size
    img = Image.new("RGB", (W, H), color)
    d = ImageDraw.Draw(img)
    d.rectangle([0, 0, W - 1, H - 1], outline=(255, 255, 255), width=12)
    d.line([(0, 0), (W, H)], fill=(255, 255, 255), width=6)
    d.line([(W, 0), (0, H)], fill=(255, 255, 255), width=6)
    ft = thai(40, True)
    bb = d.textbbox((0, 0), label, font=ft)
    d.text(((W - bb[2]) / 2, (H - bb[3]) / 2 - 5), label, font=ft, fill=(255, 255, 255))
    img.save(path, fmt)
    return path


seal = make_img(os.path.join(IMGD, "ตรา.png"), (18, 110, 160), "ตรา", (420, 420))
p1 = make_img(os.path.join(IMGD, "หลักฐาน1.jpg"), (150, 40, 130), "หลักฐาน 1", (500, 340), "JPEG")
p2 = make_img(os.path.join(IMGD, "หลักฐาน2.jpg"), (170, 90, 20), "หลักฐาน 2", (500, 340), "JPEG")
p3 = make_img(os.path.join(IMGD, "หลักฐาน3.jpg"), (20, 120, 60), "หลักฐาน 3", (500, 340), "JPEG")

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.top_margin = sec.bottom_margin = Cm(1.8)
sec.left_margin = sec.right_margin = Cm(2.0)
st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(16)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def par(text="", bold=False, size=16, after=4, center=False):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_after = Pt(after)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if text:
        r = p.add_run(text)
        r.font.name = FONT
        r.font.size = Pt(size)
        r.bold = bold
        r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    return p


def descr(shape, text, name=""):
    dp = shape._inline.docPr
    dp.set("name", name or text)
    dp.set("descr", text)


par("ส่วนที่ 4: ทดสอบรูปภาพ", bold=True, size=18, after=10)

par("4.1 รูปเดี่ยว", bold=True, after=4)
descr(par(center=True, after=10).add_run().add_picture(seal, width=Cm(4.0)), "ตราสัญลักษณ์")

par("4.2 หลายรูปในช่องเดียว (3 รูป)", bold=True, after=4)
descr(par(center=True, after=4).add_run().add_picture(p1, width=Cm(2.8)), "หลักฐานหลายรูป")
par("   (ด้านล่างควรต่อกัน 3 รูปคนละสี)", size=13, after=10)

par("4.3 รูปในตาราง", bold=True, after=4)
t = doc.add_table(rows=2, cols=3)
t.style = "Table Grid"
t.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, h in enumerate(["รายการ", "วันที่", "ภาพ"]):
    c = t.rows[0].cells[i].paragraphs[0]
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(h); r.bold = True; r.font.name = FONT; r.font.size = Pt(14)
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
for i, txt in enumerate(["ถ่ายไฟดับหมู่ 4", "28/09/2569"]):
    c = t.rows[1].cells[i].paragraphs[0]
    r = c.add_run(txt); r.font.name = FONT; r.font.size = Pt(14)
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
pc = t.rows[1].cells[2].paragraphs[0]
pc.alignment = WD_ALIGN_PARAGRAPH.CENTER
descr(pc.add_run().add_picture(p1, width=Cm(3.0)), "ภาพในตาราง")
par(after=10)

par("4.4 / 4.5 ซ่อนแถวที่ไม่มีรูป + ลบรูปที่ไม่ต้องการ", bold=True, after=4)
t2 = doc.add_table(rows=3, cols=2)
t2.style = "Table Grid"
t2.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, h in enumerate(["รายการ", "ภาพ"]):
    c = t2.rows[0].cells[i].paragraphs[0]
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(h); r.bold = True; r.font.name = FONT; r.font.size = Pt(14)
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
for ri, (name, d) in enumerate([("แถวที่ 1 (มีรูป)", "ภาพแถว1"),
                                 ("แถวที่ 2 (ไม่มีรูป)", "ภาพแถว2")], start=1):
    c0 = t2.rows[ri].cells[0].paragraphs[0]
    r = c0.add_run(name); r.font.name = FONT; r.font.size = Pt(14)
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    pc2 = t2.rows[ri].cells[1].paragraphs[0]
    pc2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    descr(pc2.add_run().add_picture(p1, width=Cm(2.4)), d)
par(after=8)
par("   แถวที่ 2 ควรหายไป และรูปที่ไม่ต้องการควรถูกลบ", size=13, after=8)
descr(par(center=True).add_run().add_picture(seal, width=Cm(2.2)), "รูปที่ไม่ต้องการ")

doc.save(TPL)
print("แม่แบบ:", os.path.basename(TPL))


def durl(path, mime):
    with open(path, "rb") as f:
        return f"data:{mime};base64," + base64.b64encode(f.read()).decode()


data = {
    "ตราสัญลักษณ์": seal,                          # จากไฟล์
    "หลักฐานหลายรูป": [p1, durl(p2, "image/jpeg"), p3],   # ไฟล์ + base64 ผสมกัน
    "ภาพในตาราง": p2,
    "ภาพแถว1": p3,
    "ภาพแถว2": None,                               # ไม่มีรูป → ต้องซ่อนแถว
}
with open(DATA, "w", encoding="utf-8") as f:
    json.dump(data, f, ensure_ascii=False, indent=2)
print("ข้อมูล  :", os.path.basename(DATA))

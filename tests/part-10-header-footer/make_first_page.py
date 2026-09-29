# -*- coding: utf-8 -*-
"""สร้างแม่แบบหน้าแรกต่างจากหน้าถัดไป"""
import os
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
TPL = os.path.join(HERE, "templates", "หน้าแรกต่างจากหน้าถัดไป.docx")
os.makedirs(os.path.dirname(TPL), exist_ok=True)

FONT = "TH Sarabun New"


def set_font(run, size=16, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run._element.rPr.rFonts.set(qn("w:cs"), FONT)


def add_field(p, code, size=14, bold=False):
    for kind, txt in (("begin", None), (None, f" {code} "), ("separate", None)):
        r = p.add_run()
        if kind:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            r._r.append(fc)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = txt
            r._r.append(it)
        set_font(r, size, bold)
    r = p.add_run("1")
    set_font(r, size, bold)
    r = p.add_run()
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "end")
    r._r.append(fc)
    set_font(r, size, bold)


def add_bottom_border(p):
    ppr = p._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    bt = OxmlElement("w:bottom")
    bt.set(qn("w:val"), "single")
    bt.set(qn("w:sz"), "6")
    bt.set(qn("w:space"), "1")
    bt.set(qn("w:color"), "000000")
    b.append(bt)
    ppr.append(b)


doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.top_margin = Cm(3.5)
sec.bottom_margin = Cm(2.5)
sec.left_margin = Cm(3.0)
sec.right_margin = Cm(2.0)
sec.header_distance = Cm(1.5)
sec.footer_distance = Cm(1.5)

# เปิด "หน้าแรกต่างจากหน้าถัดไป"
sec.different_first_page_header_footer = True

st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(16)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
st.element.rPr.rFonts.set(qn("w:cs"), FONT)
st.paragraph_format.line_spacing = 1.5


def par(text="", align=WD_ALIGN_PARAGRAPH.LEFT, bold=False, size=16, after=0):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(after)
    if text:
        set_font(p.add_run(text), size, bold)
    return p


# ---------- ส่วนหัวหน้าแรก: มีชื่อหน่วยงาน + เส้นคั่น ----------
h1 = sec.first_page_header.paragraphs[0]
h1.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(h1.add_run("{d.หน่วยงาน}"), 14, True)
add_bottom_border(h1)

# ---------- ส่วนหัวหน้าถัดไป: เหมือนกัน ----------
h2 = sec.header.paragraphs[0]
h2.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(h2.add_run("{d.หน่วยงาน}"), 14, True)
add_bottom_border(h2)

# ---------- ส่วนท้ายหน้าแรก: ไม่มีเลขหน้า ----------
f1 = sec.first_page_footer.paragraphs[0]
f1.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(f1.add_run("(หน้าแรก)"), 12, False)

# ---------- ส่วนท้ายหน้าถัดไป: มีเลขหน้า ----------
f2 = sec.footer.paragraphs[0]
f2.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(f2.add_run("หน้า "), 14, False)
add_field(f2, "PAGE", 14, False)
set_font(f2.add_run(" จาก "), 14, False)
add_field(f2, "NUMPAGES", 14, False)
set_font(f2.add_run(" หน้า"), 14, False)

# ---------- เนื้อหา ----------
par("หนังสือรับรอง", WD_ALIGN_PARAGRAPH.CENTER, True, 18, 12)
par("เรื่อง  ทดสอบหน้าแรกต่างจากหน้าถัดไป", after=14)
par("เนื้อหาหน้าแรก", bold=True, after=8)
par("{d.เนื้อหา}", after=12)

doc.add_page_break()
par("หน้าที่สอง", WD_ALIGN_PARAGRAPH.CENTER, True, 18, 14)
for i in range(1, 13):
    par(f"บรรทัดที่ {i} ของหน้าสอง — ใช้ตรวจว่าหัวเอกสารและเลขหน้าแสดงถูกต้อง", after=6)

doc.add_page_break()
par("หน้าที่สาม", WD_ALIGN_PARAGRAPH.CENTER, True, 18, 14)
for i in range(1, 13):
    par(f"บรรทัดที่ {i} ของหน้าสาม — ใช้ตรวจว่าเลขหน้าเรียงต่อเนื่อง", after=6)

doc.save(TPL)
print("สร้างแม่แบบ:", os.path.basename(TPL))
print("  - หน้าแรก : มีหัวเอกสาร, ท้ายเอกสารไม่มีเลขหน้า")
print("  - หน้าอื่น : มีหัวเอกสาร, ท้ายเอกสารมีเลขหน้า 2-3")
